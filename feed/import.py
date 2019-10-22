from contextlib import contextmanager
from csv import DictWriter
from os import environ
from pathlib import Path
from re import sub
from subprocess import run
from tempfile import TemporaryDirectory
from time import sleep
from typing import Iterator, NamedTuple, Optional

from openpyxl import load_workbook
from rows.utils import load_schema, pgimport


DATA_DIRECTORY = Path("/mnt/data/")
SCHEMA_DIRECTORY = Path("/mnt/schemas/")
CNAE_FILE = "CNAE_Subclasses_2_3_Estrutura_Detalhada.xlsx"
POSTGRES_URI = environ.get("POSTGRES_URI")

Cnae = NamedTuple("CNAE", (("codigo", int), ("descricao", str)))
Schema = NamedTuple("Schema", (("field_name", str), ("field_type", str)))


def psql(sql, capture_output=False):
    return run(["psql", POSTGRES_URI, "-c", sql], capture_output=capture_output)


def parse_code(code: str) -> Optional[int]:
    if not code:
        return

    cleaned = sub(r"\D", "", code)
    try:
        return int(cleaned)
    except ValueError:
        return


def cnaes(excel_file: Optional[str]) -> Iterator[Cnae]:
    excel_file = excel_file or str(DATA_DIRECTORY / CNAE_FILE)
    wb = load_workbook(excel_file)
    for row in wb.active.rows:
        code = parse_code(row[4].value)
        description = row[5].value

        if not all((code, description)):
            continue

        yield Cnae(code, description)


def csv_with_headers(path, rows, *fieldnames):
    with path.open("w") as fobj:
        writer = DictWriter(fobj, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(row._asdict() for row in rows)


@contextmanager
def cnaes_csv(excel_file=None):
    with TemporaryDirectory() as path:
        tmp = Path(path)
        data = tmp / "data.csv"
        schema = tmp / "schema.csv"

        csv_with_headers(data, cnaes(excel_file), "codigo", "descricao")
        csv_with_headers(
            schema,
            (Schema("codigo", "integer"), Schema("descricao", "text")),
            "field_name",
            "field_type",
        )
        yield data.as_posix(), schema.as_posix()


def create_index(table, field="cnpj"):
    psql(f"CREATE INDEX idx_{table}_{field} ON {table}({field});")


def wait_for_postgres(retries=64):
    result = psql("SELECT 1", capture_output=True)
    if not result.stderr:
        return True

    sleep(1)
    print("Waiting for PostgreSQL…")
    retries -= 1
    return wait_for_postgres(retries)


def tables_exist(retries=16):
    sql = """
        SELECT COUNT(*)
        FROM information_schema.tables
        WHERE table_schema = 'public'
    """
    result = psql(sql, capture_output=True)

    # if the count line (3rd line) is different than zero, thus tables exist
    return True if int(result.stdout.split(b"\n")[2]) != 0 else False


def main():
    if wait_for_postgres() and tables_exist():
        print(
            (
                "There are existing tables in the database. "
                "Please, start with a clean database."
            )
        )
        return

    tables = ("empresa", "socio", "cnae_secundaria")
    files = ("empresa", "socio", "cnae-secundaria")
    for table, filename in zip(tables, files):
        pgimport(
            str(DATA_DIRECTORY / f"{filename}.csv.gz"),
            POSTGRES_URI,
            table,
            schema=load_schema(str(SCHEMA_DIRECTORY / f"{filename}.csv")),
        )
        create_index(table)

    with cnaes_csv() as (source, schema):
        pgimport(source, POSTGRES_URI, "cnae", schema=load_schema(schema))
        create_index("cnae", "codigo")


if __name__ == "__main__":
    main()
